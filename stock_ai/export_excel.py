"""
分析結果を Excel ファイルに出力する。

シート構成:
  1. 注目銘柄ランキング  - 総合スコア上位銘柄
  2. 全銘柄分析結果      - フィルター前全銘柄
  3. テクニカル詳細      - テクニカル指標詳細
  4. 出来高資金流入      - 出来高・売買代金詳細
  5. 決算モメンタム      - 決算データ（将来対応）
  6. ファンダメンタル    - 財務データ（将来対応）
  7. 市場指数            - 日経・NASDAQ 等
  8. 実行ログ            - 実行情報サマリー
"""

import logging
import sqlite3
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import DB_PATH, EXCEL_DIR, MARKET_SYMBOLS

logger = logging.getLogger(__name__)

# 配色
_C = {
    "rank":    "1F497D",  # 紺
    "all":     "375623",  # 緑
    "tech":    "C55A11",  # オレンジ
    "vol":     "7030A0",  # 紫
    "earn":    "833C00",  # 茶
    "fund":    "0070C0",  # 青
    "idx":     "4472C4",  # 水色
    "log":     "4A4A4A",  # グレー
    "intraday": "C00000", # 赤
    "profit":  "548235",  # 緑
    "signal":  "B45F06",  # 濃いオレンジ
    "report":  "203864",  # 濃紺
    "text":    "FFFFFF",
}


# ── ユーティリティ ────────────────────────────────────────────


def _style_header(ws, color: str, col_count: int):
    fill = PatternFill("solid", fgColor=color)
    font = Font(bold=True, color=_C["text"])
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22


def _auto_width(ws, max_w: int = 40):
    for col in ws.columns:
        w = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 4, max_w)


def _fmt(val):
    if val is None:
        return ""
    if isinstance(val, float):
        return round(val) if val > 1_000_000 else round(val, 2)
    return val


def _write_df(ws, df: pd.DataFrame, col_map: dict, color: str):
    """col_map = {df_col: 表示名} で列を選択してシートに書き込む"""
    ws.freeze_panes = "A2"
    ws.append(list(col_map.values()))
    _style_header(ws, color, len(col_map))
    for _, row in df.iterrows():
        ws.append([_fmt(row.get(k)) for k in col_map])
    _auto_width(ws)


# ── DB からのデータロード ──────────────────────────────────────


def _load_analysis_results(target_date: str) -> pd.DataFrame:
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql_query(
        "SELECT * FROM analysis_results WHERE date = ? ORDER BY rank NULLS LAST",
        conn, params=(target_date,)
    )
    conn.close()
    return df


def _load_all_quotes(target_date: str) -> pd.DataFrame:
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql_query(
        """
        SELECT dq.code, lc.company_name, dq.close,
               dq.open, dq.high, dq.low, dq.volume, dq.turnover_value
        FROM daily_quotes dq
        LEFT JOIN listed_companies lc ON dq.code = lc.code
        WHERE dq.date = ?
        ORDER BY dq.turnover_value DESC NULLS LAST
        """,
        conn, params=(target_date,)
    )
    conn.close()
    return df


def _load_market_indices() -> pd.DataFrame:
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql_query(
        "SELECT * FROM market_indices ORDER BY symbol, date DESC",
        conn
    )
    conn.close()
    df["name"] = df["symbol"].map(MARKET_SYMBOLS)
    return df


def _load_earnings() -> pd.DataFrame:
    conn = sqlite3.connect(DB_PATH)
    try:
        df = pd.read_sql_query(
            "SELECT * FROM earnings_data ORDER BY code, fiscal_period DESC",
            conn
        )
    except Exception:
        df = pd.DataFrame()
    conn.close()
    return df


def _load_fundamentals() -> pd.DataFrame:
    conn = sqlite3.connect(DB_PATH)
    try:
        df = pd.read_sql_query(
            "SELECT * FROM fundamentals ORDER BY code",
            conn
        )
    except Exception:
        df = pd.DataFrame()
    conn.close()
    return df


# ── 各シート書き込み ──────────────────────────────────────────


def _write_ranking(wb, df: pd.DataFrame):
    ws = wb.create_sheet("注目銘柄ランキング")
    col_map = {
        "rank":                    "順位",
        "code":                    "銘柄コード",
        "company_name":            "銘柄名",
        "close":                   "終値",
        "change_pct":              "前日比(%)",
        "volume":                  "出来高",
        "volume_ratio_5d":         "出来高倍率(5日)",
        "trading_value":           "売買代金",
        "technical_score":         "テクニカルスコア",
        "volume_flow_score":       "出来高・資金流入スコア",
        "earnings_momentum_score": "決算モメンタムスコア",
        "fundamental_score":       "ファンダメンタルスコア",
        "total_score":             "総合スコア",
        "reason":                  "選定理由",
    }
    _write_df(ws, df, col_map, _C["rank"])
    logger.info("注目銘柄ランキングシート: %d 行", len(df))


def _write_all_results(wb, df: pd.DataFrame):
    ws = wb.create_sheet("全銘柄分析結果")
    col_map = {
        "code":          "銘柄コード",
        "company_name":  "銘柄名",
        "close":         "終値",
        "open":          "始値",
        "high":          "高値",
        "low":           "安値",
        "volume":        "出来高",
        "turnover_value": "売買代金",
    }
    _write_df(ws, df, col_map, _C["all"])
    logger.info("全銘柄分析結果シート: %d 行", len(df))


def _write_technical(wb, df: pd.DataFrame):
    ws = wb.create_sheet("テクニカル詳細")
    col_map = {
        "rank":           "順位",
        "code":           "銘柄コード",
        "company_name":   "銘柄名",
        "close":          "終値",
        "change_pct":     "前日比(%)",
        "ma5":            "MA5",
        "ma25":           "MA25",
        "ma5_gap_pct":    "MA5乖離(%)",
        "ma25_gap_pct":   "MA25乖離(%)",
        "high_20d":       "20日高値",
        "high_breakout":  "20日高値更新",
        "trend_5d":       "5日上昇日数",
        "technical_score": "テクニカルスコア",
    }
    _write_df(ws, df, col_map, _C["tech"])


def _write_volume_flow(wb, df: pd.DataFrame):
    ws = wb.create_sheet("出来高資金流入")
    col_map = {
        "rank":                   "順位",
        "code":                   "銘柄コード",
        "company_name":           "銘柄名",
        "close":                  "終値",
        "volume":                 "出来高",
        "volume_ma5":             "5日平均出来高",
        "volume_ratio_5d":        "出来高倍率(5日)",
        "trading_value":          "売買代金",
        "trading_value_ma5":      "5日平均売買代金",
        "trading_value_ratio_5d": "売買代金倍率(5日)",
        "volume_flow_score":      "出来高・資金流入スコア",
    }
    _write_df(ws, df, col_map, _C["vol"])


def _write_earnings(wb, df: pd.DataFrame):
    ws = wb.create_sheet("決算モメンタム")
    if df.empty:
        ws.append(["（決算データなし。J-Quants有料プランまたはCSV取込後に表示されます）"])
        _style_header(ws, _C["earn"], 1)
        return

    col_map = {
        "code":                       "銘柄コード",
        "fiscal_period":              "決算期",
        "sales_growth_pct":           "売上成長率(%)",
        "operating_profit_growth_pct": "営業利益成長率(%)",
        "ordinary_profit_growth_pct": "経常利益成長率(%)",
        "net_income_growth_pct":      "純利益成長率(%)",
        "eps_growth_pct":             "EPS成長率(%)",
        "progress_rate_pct":          "進捗率(%)",
        "upward_revision_flag":       "上方修正",
    }
    _write_df(ws, df, col_map, _C["earn"])


def _write_fundamentals(wb, df: pd.DataFrame):
    ws = wb.create_sheet("ファンダメンタル")
    if df.empty:
        ws.append(["（財務データなし。J-Quants有料プラン・Kabutan・IR BANK CSV取込後に表示されます）"])
        _style_header(ws, _C["fund"], 1)
        return

    col_map = {
        "code":             "銘柄コード",
        "market_cap":       "時価総額",
        "per":              "PER",
        "pbr":              "PBR",
        "roe":              "ROE(%)",
        "equity_ratio":     "自己資本比率(%)",
        "operating_margin": "営業利益率(%)",
        "sales":            "売上高",
        "operating_profit": "営業利益",
        "eps":              "EPS",
        "dividend_yield":   "配当利回り(%)",
    }
    _write_df(ws, df, col_map, _C["fund"])


def _write_indices(wb, df: pd.DataFrame):
    ws = wb.create_sheet("市場指数")
    col_map = {
        "symbol": "シンボル",
        "name":   "名称",
        "date":   "日付",
        "open":   "始値",
        "high":   "高値",
        "low":    "安値",
        "close":  "終値",
        "volume": "出来高",
    }
    _write_df(ws, df, col_map, _C["idx"])


def _write_log(wb, target_date: str, ranked: int, all_quotes: int):
    ws = wb.create_sheet("実行ログ")
    rows = [
        ("実行日時",       datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("分析対象日",     target_date),
        ("注目銘柄数",     ranked),
        ("全銘柄取得数",   all_quotes),
        ("株価条件",       "50〜3000円"),
        ("前日比条件",     "+3%以上"),
        ("出来高倍率条件", "5日平均比2倍以上"),
        ("売買代金条件",   "5000万円以上"),
    ]
    ws.append(["項目", "値"])
    _style_header(ws, _C["log"], 2)
    for label, val in rows:
        ws.append([label, val])
    _auto_width(ws)


def _write_intraday_prices(wb, df: pd.DataFrame):
    ws = wb.create_sheet("5分足データ")
    col_map = {
        "code":     "銘柄コード",
        "datetime": "日時",
        "open":     "始値",
        "high":     "高値",
        "low":      "安値",
        "close":    "終値",
        "volume":   "出来高",
    }
    _write_df(ws, df, col_map, _C["intraday"])
    logger.info("5分足データシート: %d 行", len(df))


def _write_profit_calc(wb, df: pd.DataFrame):
    ws = wb.create_sheet("損益計算")
    col_map = {
        "code":          "銘柄コード",
        "entry_price":   "エントリー価格",
        "current_price": "現在価格",
        "profit_pct":    "損益率(%)",
        "entry_mode":    "エントリーモード",
        "calculated_at": "計算時刻",
    }
    _write_df(ws, df, col_map, _C["profit"])
    logger.info("損益計算シート: %d 行", len(df))


def _write_trade_signals(wb, df: pd.DataFrame):
    ws = wb.create_sheet("デイトレ判定")
    col_map = {
        "code":                "銘柄コード",
        "signal_datetime":     "判定時刻",
        "current_price":       "現在価格",
        "entry_price":         "エントリー価格",
        "profit_pct":          "損益率(%)",
        "signal":              "シグナル",
        "reason":              "理由",
        "signal_score":        "シグナルスコア",
        "risk_level":          "リスクレベル",
        "positive_factors":    "プラス材料",
        "negative_factors":    "マイナス材料",
        "vwap":                "VWAP",
        "volume_ma3":          "出来高3本平均",
        "volume_ma6":          "出来高6本平均",
        "volume_decline_flag": "出来高減少フラグ",
        "bar_change_pct":      "直近1本変化率(%)",
        "abnormal_volume_flag": "出来高異常フラグ",
        "confirmation_count":  "連続確認本数",
        "signal_strength":     "シグナル強度",
        "previous_high":              "前日高値",
        "previous_low":               "前日安値",
        "previous_close":             "前日終値",
        "breakout_prev_high_flag":    "前日高値ブレイクフラグ",
        "breakdown_prev_low_flag":    "前日安値割れフラグ",
        "opening_30min_high":         "寄り付き30分高値",
        "opening_30min_low":          "寄り付き30分安値",
        "opening_range_breakout":     "寄り付き高値ブレイクフラグ",
        "opening_range_breakdown":    "寄り付き安値割れフラグ",
        "volume_ma12":                "出来高12本平均",
        "volume_ratio_3_12":          "出来高比率(3本/12本)",
        "volume_surge_continuation":  "出来高急増継続フラグ",
        "volume_fading":              "出来高失速フラグ",
        "atr_14":                     "ATR(14)",
        "atr_stop_price":             "ATR損切りライン",
        "atr_stop_loss_flag":         "ATR損切りフラグ",
        "signal_changed":             "シグナル変化フラグ",
        "market_sentiment":           "市場地合い",
        "market_change_pct":          "市場平均前日比(%)",
        "entry_score":                "ENTRY_SCORE",
        "entry_candidate":            "買いエントリー候補判定",
        "entry_factors":              "ENTRY_SCORE内訳",
    }
    _write_df(ws, df, col_map, _C["signal"])
    logger.info("デイトレ判定シート: %d 行", len(df))


def _write_signal_history(wb, df: pd.DataFrame):
    ws = wb.create_sheet("シグナル変化履歴")
    col_map = {
        "code":             "銘柄コード",
        "previous_signal":  "前回シグナル",
        "current_signal":   "現在シグナル",
        "changed_flag":     "変化フラグ",
        "reason":           "理由",
        "signal_datetime":  "判定時刻",
    }
    _write_df(ws, df, col_map, _C["signal"])
    logger.info("シグナル変化履歴シート: %d 行", len(df))


def _write_entry_candidate_history(wb, df: pd.DataFrame):
    ws = wb.create_sheet("買い候補変化履歴")
    col_map = {
        "code":                      "銘柄コード",
        "previous_entry_candidate":  "前回買い候補判定",
        "current_entry_candidate":   "現在買い候補判定",
        "changed_flag":              "変化フラグ",
        "entry_score":               "ENTRY_SCORE",
        "entry_factors":             "ENTRY_SCORE内訳",
        "reason":                    "理由",
        "signal_datetime":           "判定時刻",
    }
    _write_df(ws, df, col_map, _C["profit"])
    logger.info("買い候補変化履歴シート: %d 行", len(df))


def _write_daily_report(wb, df: pd.DataFrame):
    ws = wb.create_sheet("日次監視レポート")
    col_map = {
        "code":              "銘柄コード",
        "entry_count":       "ENTRY回数",
        "watch_count":        "WATCH回数",
        "take_profit_count": "TAKE_PROFIT回数",
        "stop_loss_count":   "STOP_LOSS回数",
        "max_profit_pct":    "最大含み益(%)",
        "max_loss_pct":      "最大含み損(%)",
        "last_signal":       "最終シグナル",
        "last_reason":       "判定理由",
    }
    _write_df(ws, df, col_map, _C["report"])
    logger.info("日次監視レポートシート: %d 行", len(df))


# ── エントリーポイント ─────────────────────────────────────────


def export_to_excel(df_ranked: pd.DataFrame, target_date: str) -> Path:
    """
    分析結果を Excel に出力して保存パスを返す。
    df_ranked が空でも空ファイルとして出力する。
    """
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    df_all      = _load_all_quotes(target_date)
    df_idx      = _load_market_indices()
    df_earn     = _load_earnings()
    df_fund     = _load_fundamentals()

    _write_ranking(wb, df_ranked)
    _write_all_results(wb, df_all)
    _write_technical(wb, df_ranked)
    _write_volume_flow(wb, df_ranked)
    _write_earnings(wb, df_earn)
    _write_fundamentals(wb, df_fund)
    _write_indices(wb, df_idx)
    _write_log(wb, target_date, len(df_ranked), len(df_all))

    out_path = EXCEL_DIR / f"stock_analysis_{target_date}.xlsx"
    wb.save(out_path)
    logger.info("Excel 出力完了: %s", out_path)
    return out_path


def export_intraday_excel(
    df_prices: pd.DataFrame,
    run_dt: str,
    df_positions: pd.DataFrame | None = None,
    df_signals: pd.DataFrame | None = None,
    df_history: pd.DataFrame | None = None,
    df_entry_history: pd.DataFrame | None = None,
) -> Path:
    """
    5分足データ（と、あれば損益計算結果・デイトレ判定結果・シグナル変化履歴・
    買い候補変化履歴）を Excel に出力して保存パスを返す。
    """
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_intraday_prices(wb, df_prices)
    if df_positions is not None and not df_positions.empty:
        _write_profit_calc(wb, df_positions)
    if df_signals is not None and not df_signals.empty:
        _write_trade_signals(wb, df_signals)
    if df_history is not None and not df_history.empty:
        _write_signal_history(wb, df_history)
    if df_entry_history is not None and not df_entry_history.empty:
        _write_entry_candidate_history(wb, df_entry_history)

    out_path = EXCEL_DIR / f"intraday_prices_{run_dt}.xlsx"
    wb.save(out_path)
    logger.info("Excel 出力完了: %s", out_path)
    return out_path


def export_daily_report_excel(df_report: pd.DataFrame, target_date: str) -> Path:
    """
    日次監視レポートを Excel に出力して保存パスを返す。
    """
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_daily_report(wb, df_report)

    out_path = EXCEL_DIR / f"daily_report_{target_date}.xlsx"
    wb.save(out_path)
    logger.info("Excel 出力完了: %s", out_path)
    return out_path


def _write_backtest_summary(wb, summary: dict):
    ws = wb.create_sheet("バックテスト結果")
    rows = [
        ("総トレード数",         summary["total_trades"]),
        ("勝率(%)",              summary["win_rate_pct"]),
        ("平均利益率(%)",        summary["avg_profit_pct"]),
        ("平均損失率(%)",        summary["avg_loss_pct"]),
        ("最大ドローダウン(%)",  summary["max_drawdown_pct"]),
        ("TAKE_PROFIT到達率(%)", summary["take_profit_rate_pct"]),
        ("STOP_LOSS到達率(%)",   summary["stop_loss_rate_pct"]),
    ]
    ws.append(["項目", "値"])
    _style_header(ws, _C["report"], 2)
    for label, val in rows:
        ws.append([label, val])
    _auto_width(ws)
    logger.info("バックテスト結果シート出力完了")


def _write_backtest_trades(wb, df: pd.DataFrame):
    ws = wb.create_sheet("バックテスト詳細")
    col_map = {
        "code":              "銘柄コード",
        "date":              "対象日",
        "exit_reason":       "決済理由",
        "profit_pct":        "損益率(%)",
        "max_drawdown_pct":  "最大ドローダウン(%)",
    }
    _write_df(ws, df, col_map, _C["report"])
    logger.info("バックテスト詳細シート: %d 行", len(df))


def export_backtest_excel(summary: dict, df_trades: pd.DataFrame, run_dt: str) -> Path:
    """
    バックテスト結果（サマリー・トレード詳細）を Excel に出力して保存パスを返す。
    """
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_backtest_summary(wb, summary)
    _write_backtest_trades(wb, df_trades)

    out_path = EXCEL_DIR / f"backtest_{run_dt}.xlsx"
    wb.save(out_path)
    logger.info("Excel 出力完了: %s", out_path)
    return out_path

